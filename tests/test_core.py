from pathlib import Path
from urllib.parse import unquote
import asyncio
import io
import json
import shutil
import zipfile
import pytest
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
from docx import Document
from server.storage import Store,Conflict
from server.importers import parse_file,stage,validate_public_url
from server.research import ResearchManager,validate_draft
from server.exports import backup,inspect_backup,preview_restore,restore,export_html,markdown_bundle
from server.app import create_app

ROOT=Path(__file__).resolve().parents[1]
BUILTIN=ROOT/'resources/builtin'

@pytest.fixture(scope='session')
def seed(tmp_path_factory):
    directory=tmp_path_factory.mktemp('seed')
    Store(directory,BUILTIN)
    return directory

@pytest.fixture
def store(tmp_path,seed):
    shutil.copy2(seed/'personal.sqlite',tmp_path/'personal.sqlite')
    return Store(tmp_path,BUILTIN)

def library(store,title):
    p=store.propose('library_create',{'title':title})
    return store.confirm(p['id'],p['hash'])['id']

def add_text(store,lid,text='太阳病，头痛发热，汗出恶风。',name='测试.txt'):
    raw=text.encode()
    p=stage(store,raw,parse_file(name,raw),[lid],name)
    return store.confirm(p['id'],p['hash'])['id']

def test_baseline_readonly_and_apis(store):
    app=create_app(store.root,BUILTIN,ROOT/'ui','test-token')
    with TestClient(app,headers={'x-shanghan-token':'test-token'}) as client:
        meta=client.get('/api/meta').json()
        assert meta['passage_count']==398 and len(meta['commentators'])==9
        assert len(meta['groups'])==23 and meta['ai'] is False
        assert len(client.get('/api/origbooks').json())==9
        assert client.get('/api/passage/12').json()['textbook']
        assert client.get('/api/search?q=太阳').json()['count']>0
        assert client.get('/').status_code==200
        assert client.get('/api/v1/status',headers={'x-shanghan-token':'wrong'}).status_code==401
    assert len(store.sources(include_builtin=True))==20

def test_proposal_cancel_replay_conflict(store):
    p=store.propose('library_create',{'title':'取消专题'})
    store.cancel(p['id'])
    with pytest.raises(Conflict):store.confirm(p['id'],p['hash'])
    assert not store.libraries()
    p=store.propose('library_create',{'title':'桂枝汤证'})
    first=store.confirm(p['id'],p['hash'])
    assert first==store.confirm(p['id'],p['hash'])
    assert len(store.libraries())==1
    stale=store.propose('library_update',{'id':first['id'],'title':'旧预览'})
    current=store.propose('library_update',{'id':first['id'],'title':'新名称'})
    store.confirm(current['id'],current['hash'])
    with pytest.raises(Conflict):store.confirm(stale['id'],stale['hash'])

def test_scope_duplicate_and_historical_reference(store):
    a,b=library(store,'太阳'),library(store,'太阴')
    sid=add_text(store,a)
    other=add_text(store,b,'太阴之为病，腹满而吐。','太阴.txt')
    scope=store.snapshot([a])
    assert len(store.search('太阳',scope))==1
    assert store.search('太阴',scope)==[]
    ch=store.search('太阳',scope)[0]
    with pytest.raises(ValueError):store.chunk(ch['id'],store.snapshot([b]))
    assert add_text(store,b)==sid
    assert len(store.sources())==2
    p=store.propose('source_delete',{'id':sid})
    store.confirm(p['id'],p['hash'])
    assert store.chunk(ch['id'],scope)['text']==ch['text']
    assert not store.sources(a)
    assert other in store.snapshot([b])['source_ids']

def test_builtin_search_variants(store):
    scope=store.snapshot([],True)
    assert store.search('太陽',scope)
    assert store.search('桂枝汤',scope)
    sid=store.sources(include_builtin=True)[0]['id']
    with pytest.raises(ValueError):store.propose('source_delete',{'id':sid})

def test_empty_import_not_committed(store):
    lid=library(store,'扫描')
    writer=PdfWriter();writer.add_blank_page(200,200)
    data=io.BytesIO();writer.write(data);raw=data.getvalue()
    parsed=parse_file('scan.pdf',raw)
    assert parsed['metadata']['pages_without_text']==[1]
    p=stage(store,raw,parsed,[lid],'scan.pdf')
    with pytest.raises(ValueError):store.confirm(p['id'],p['hash'])
    assert store.sources()==[]

def pdf_sample(mixed=False):
    writer=PdfWriter();page=writer.add_blank_page(400,400)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
    stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 30 300 Td (Guizhi Tang source text) Tj ET')
    page[NameObject('/Contents')]=writer._add_object(stream)
    if mixed:writer.add_blank_page(400,400)
    data=io.BytesIO();writer.write(data);return data.getvalue()

def test_supported_parsers():
    assert parse_file('test.txt','太阳病'.encode('gb18030'))['chunks'][0]['text']=='太阳病'
    assert parse_file('test.md','# 桂枝汤\n\n太阳中风。'.encode())['chunks'][-1]['locator'].startswith('桂枝汤')
    doc=Document();doc.add_heading('桂枝汤',1);doc.add_paragraph('太阳中风，汗出恶风。');table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='共识';table.cell(0,1).text='证据'
    data=io.BytesIO();doc.save(data)
    assert any('共识 | 证据' in c['text'] for c in parse_file('test.docx',data.getvalue())['chunks'])
    pdf=parse_file('test.pdf',pdf_sample(True))
    assert pdf['metadata']['pages_without_text']==[2]
    assert pdf['chunks'][0]['locator'].startswith('文件第1页')
    with pytest.raises(ValueError):parse_file('legacy.doc',b'test')
    data=io.BytesIO()
    with zipfile.ZipFile(data,'w') as z:
        z.writestr('META-INF/container.xml','<container><rootfiles><rootfile full-path="OEBPS/book.opf"/></rootfiles></container>')
        z.writestr('OEBPS/book.opf','<package xmlns="http://www.idpf.org/2007/opf"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>测试书籍</dc:title></metadata><manifest><item id="one" href="one.xhtml"/></manifest><spine><itemref idref="one"/></spine></package>')
        z.writestr('OEBPS/one.xhtml','<html><body><h1>桂枝汤</h1><p>太阳中风</p></body></html>')
    result=parse_file('test.epub',data.getvalue())
    assert result['title']=='测试书籍' and result['chunks']

def test_quote_validation():
    evidence=[{'evidence_id':'E001','id':'s:0','source_id':'s','text':'太阳病，汗出恶风。','title':'原书','author':'注家','locator':'段1','quote':'太阳病'}]
    data={'comparison':[{'text':'有效','citations':[{'evidence_id':'E001','quote':'汗出恶风'}]},{'text':'伪造','citations':[{'evidence_id':'E001','quote':'无汗身痛'}]}],'gaps':['其他资料未覆盖']}
    draft=validate_draft('问题',data,evidence)
    assert '有效' in draft['body'] and '伪造' not in draft['body']
    assert draft['warnings'] and len(draft['evidence'])==1
    assert '<script>' not in export_html({'title':'测试','body':'<script>alert(1)</script>','evidence':[]})

def test_backup_roundtrip_tamper(store):
    lid=library(store,'备份');sid=add_text(store,lid)
    p=store.propose('note_save',{'title':'笔记','body':'已核对原文','evidence':[]})
    store.confirm(p['id'],p['hash'])
    result=backup(store);raw=Path(result['path']).read_bytes()
    manifest=inspect_backup(raw)
    assert all('credential' not in name for name in manifest['files'])
    preview=preview_restore(store,raw)
    assert preview['contents']=={'libraries':['备份'],'sources':['测试'],'notes':['笔记'],
                                 'projects':[],'archived_libraries':[],'archived_projects':[]}
    restored=restore(store,preview['id'],preview['hash'])
    target=Store(restored['directory'],Path(restored['directory'])/'builtin')
    assert target.sources()[0]['id']==sid and target.notes()[0]['title']=='笔记'
    assert target.search('太阳',target.snapshot([lid]))
    bad=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as z,zipfile.ZipFile(bad,'w') as out:
        for name in z.namelist():out.writestr(name,b'bad' if name=='personal.sqlite' else z.read(name))
    with pytest.raises(ValueError):inspect_backup(bad.getvalue())

def test_markdown_relative_attachments(store):
    lid=library(store,'导出');sid=add_text(store,lid)
    chunk=store.search('太阳',store.snapshot([lid]))[0]
    note={'title':'导出验证','body':'## 比较\n\n原文可核对。','evidence':[{**chunk,'quote':'汗出恶风'}]}
    result=markdown_bundle(store,note)
    target=Path(result['markdown_path'])
    body=unquote(target.read_text(encoding='utf-8'))
    snapshot=next(Path(result['attachments_path']).glob('来源-*.md'))
    original=next(Path(result['attachments_path']).glob('原件-*'))
    assert snapshot.relative_to(target.parent).as_posix() in body
    assert original.read_bytes()==(store.root/store.source(sid)['original']).read_bytes()
    assert original.name in unquote(snapshot.read_text(encoding='utf-8'))
    assert chunk['text'] in snapshot.read_text(encoding='utf-8')
    assert str(store.root) not in body
    builtin=store.search('桂枝汤',store.snapshot([],True))[0]
    result=markdown_bundle(store,{'title':'经典','body':'原文','evidence':[{**builtin,'quote':'桂枝汤'}]})
    assert len(list(Path(result['attachments_path']).iterdir()))==1
    note['evidence'][0]['quote']='不存在的伪造引文'
    with pytest.raises(ValueError):markdown_bundle(store,note)

def test_import_original_change_rolls_back(store):
    lid=library(store,'回滚');raw='原文'.encode()
    p=stage(store,raw,parse_file('x.txt',raw),[lid],'x.txt')
    (store.root/p['payload']['original']).write_bytes(b'changed')
    with pytest.raises(Conflict):store.confirm(p['id'],p['hash'])
    assert not store.sources()

class ScriptedModel:
    def __init__(self,cid):self.calls=0;self.cid=cid
    def config(self):return {'configured':True}
    async def complete(self,messages,max_tokens=4000):
        self.calls+=1
        if self.calls==1:return json.dumps({'queries':['太阳'],'read_ids':[],'finish':False}),{}
        if self.calls==2:return json.dumps({'queries':[],'read_ids':[self.cid],'finish':True}),{}
        return json.dumps({'comparison':[{'text':'测试观点','citations':[{'evidence_id':'E001','quote':'汗出恶风'}]}],'gaps':['这是确定性测试，非真实学术结论']}),{}

@pytest.mark.asyncio
async def test_research_end_to_end_and_stop(store):
    lid=library(store,'研究');add_text(store,lid)
    scope=store.snapshot([lid]);cid=store.search('太阳',scope)[0]['id']
    manager=ResearchManager(store,ScriptedModel(cid))
    task=manager.start('比较桂枝汤证解释',scope)
    await manager.running[task['id']]
    result=store.task(task['id'])
    assert result['status']=='awaiting_review' and result['draft']['evidence']
    proposal=manager.propose_save(task['id']);store.confirm(proposal['id'],proposal['hash'])
    assert len(store.notes())==1
    repeated=manager.propose_save(task['id']);assert repeated['kind']=='note_update'
    store.confirm(repeated['id'],repeated['hash']);assert len(store.notes())==1
    class Slow(ScriptedModel):
        async def complete(self,*args,**kwargs):await asyncio.sleep(60)
    manager=ResearchManager(store,Slow(cid))
    task=manager.start('停止测试',scope);await asyncio.sleep(.02)
    await manager.stop(task['id'])
    assert store.task(task['id'])['status']=='stopped'

def test_web_rejects_local_addresses():
    with pytest.raises(ValueError):validate_public_url('http://127.0.0.1/private')
    with pytest.raises(ValueError):validate_public_url('file:///C:/secret')

@pytest.mark.asyncio
async def test_research_beyond_eight_failure_resume_and_restart(store):
    lid=library(store,'限额');add_text(store,lid)
    scope=store.snapshot([lid]);cid=store.search('太阳',scope)[0]['id']
    class Endless(ScriptedModel):
        count = 0
        async def complete(self,messages,*args,**kwargs):
            if 'evidence' in json.loads(messages[-1]['content']):
                return json.dumps({'general_answer':'长流程完成后的模型回答，未作为原文结论。'}),{}
            self.count += 1
            return json.dumps({'queries':['太阳'],'read_ids':[],'finish':self.count == 12}),{}
    manager=ResearchManager(store,Endless(cid))
    task=manager.start('有限循环',scope);await manager.running[task['id']]
    assert store.task(task['id'])['round']==12 and store.task(task['id'])['status']=='awaiting_review'
    manager.gateway=ScriptedModel(cid)
    manager.resume(task['id']);await manager.running[task['id']]
    assert store.task(task['id'])['status']=='awaiting_review'
    class Broken(ScriptedModel):
        async def complete(self,*args,**kwargs):raise ValueError('模拟服务失败')
    manager.gateway=Broken(cid)
    failure=manager.start('失败恢复',scope);await manager.running[failure['id']]
    assert store.task(failure['id'])['status']=='failed'
    store.update_task(failure['id'],status='running')
    restarted=Store(store.root,BUILTIN)
    assert restarted.task(failure['id'])['status']=='interrupted'

@pytest.mark.asyncio
async def test_research_rejects_model_guessed_out_of_scope_id(store):
    a,b=library(store,'A'),library(store,'B');add_text(store,a);add_text(store,b,'不得越库读取的内容')
    outside=store.search('越库',store.snapshot([b]))[0]['id']
    class Malicious(ScriptedModel):
        async def complete(self,*args,**kwargs):return json.dumps({'queries':[],'read_ids':[outside],'finish':True}),{}
    manager=ResearchManager(store,Malicious(outside));task=manager.start('范围测试',store.snapshot([a]))
    await manager.running[task['id']]
    assert store.task(task['id'])['status']=='failed' and not store.task(task['id'])['evidence']

def test_same_author_builtin_not_two_independent_sources():
    evidence=[{'evidence_id':f'E00{i}','id':str(i),'source_id':'builtin:version:'+str(i),'text':'汗出恶风','title':title,'author':'同一注家','locator':'段1'} for i,title in enumerate(['原书','逐条注文'],1)]
    draft=validate_draft('比较',{'consensus':[{'text':'不能误算成共识','citations':[{'evidence_id':'E001','quote':'汗出恶风'},{'evidence_id':'E002','quote':'汗出恶风'}]}]},evidence)
    assert not draft['sections']['consensus'] and '不能误算成共识' not in draft['body']

@pytest.mark.asyncio
async def test_web_parser_success_partial_and_failure(monkeypatch):
    import server.importers as module
    import httpx
    real_client=httpx.AsyncClient
    monkeypatch.setattr(module,'validate_public_url',lambda url:None)
    def handler(request):
        if request.url.path=='/article':return httpx.Response(200,headers={'content-type':'text/html'},text='<html><title>公开文章测试</title><article><h1>文章正文</h1><p>'+('资料正文，太阳病，汗出恶风。'*12)+'</p></article></html>')
        if request.url.path=='/generic':return httpx.Response(200,headers={'content-type':'text/html'},text='<html><body>'+('资料正文。'*30)+'</body></html>')
        return httpx.Response(403)
    monkeypatch.setattr(module.httpx,'AsyncClient',lambda **kwargs:real_client(transport=httpx.MockTransport(handler),**kwargs))
    raw,parsed=await module.parse_url('https://example.org/article')
    assert parsed['title']=='公开文章测试' and parsed['chunks'] and parsed['origin'].endswith('/article')
    _,partial=await module.parse_url('https://example.org/generic')
    assert partial['metadata']['warnings']
    with pytest.raises(ValueError):await module.parse_url('https://example.org/blocked')
