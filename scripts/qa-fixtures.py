from pathlib import Path
import runpy
import json
import zipfile
import sys
from docx import Document
root=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(root))
out=root/'.test-data/fixtures'
out.mkdir(parents=True,exist_ok=True)
text='程序验证样例，非学术结论。\n\n太阳病，头痛发热，汗出恶风。\n\n用于检查桂枝汤证专题的引用与范围。'
(out/'桂枝汤证_程序验证.md').write_text('# 程序验证资料\n\n'+text,encoding='utf-8')
(out/'太阴病_程序验证.txt').write_text('程序验证样例。太阴之为病，腹满而吐。',encoding='utf-8')
doc=Document();doc.add_heading('Word 程序验证',0);doc.add_paragraph(text);doc.save(out/'程序验证.docx')
pdf_sample=runpy.run_path(str(root/'tests/test_core.py'))['pdf_sample']
(out/'文字与空白混合_程序验证.pdf').write_bytes(pdf_sample(True))
with zipfile.ZipFile(out/'程序验证.epub','w') as z:
    z.writestr('META-INF/container.xml','<container><rootfiles><rootfile full-path="book.opf"/></rootfiles></container>')
    z.writestr('book.opf','<package xmlns="http://www.idpf.org/2007/opf"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>EPUB 程序验证</dc:title></metadata><manifest><item id="one" href="one.xhtml"/></manifest><spine><itemref idref="one"/></spine></package>')
    z.writestr('one.xhtml','<html><body><h1>程序验证</h1><p>'+text+'</p></body></html>')
print(out)
